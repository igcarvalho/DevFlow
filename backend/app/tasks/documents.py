import io
import uuid

from pypdf import PdfReader
from pptx import Presentation
from sqlalchemy.orm import Session

from app.db.session import SessionLocal
from app.models import Document, DocumentStatus, DocumentVersion  # noqa: F401
from app.services.storage import download_file
from app.tasks.celery_app import celery_app


def extract_text_from_pdf(file_data: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_data))
    texts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text)
    return "\n\n".join(texts)


def extract_text_from_pptx(file_data: bytes) -> str:
    presentation = Presentation(io.BytesIO(file_data))
    texts = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        slide_texts.append(text.strip())
        if slide_texts:
            texts.append(f"--- Slide {slide_number} ---\n" + "\n".join(slide_texts))
    return "\n\n".join(texts)


@celery_app.task(bind=True, name="app.tasks.documents.process_document_version")
def process_document_version(self, version_id: str) -> dict:
    db: Session = SessionLocal()
    try:
        version_uuid = uuid.UUID(version_id)
        version = db.query(DocumentVersion).filter(DocumentVersion.id == version_uuid).first()
        if version is None:
            return {"status": "error", "message": "Version not found"}

        document = db.query(Document).filter(Document.id == version.document_id).first()
        if document is None:
            return {"status": "error", "message": "Document not found"}

        document.status = DocumentStatus.processing
        db.commit()

        try:
            file_data = download_file(version.file_key)

            if version.mime_type == "application/pdf":
                text = extract_text_from_pdf(file_data)
            elif version.mime_type in (
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "application/vnd.ms-powerpoint",
            ):
                text = extract_text_from_pptx(file_data)
            else:
                raise ValueError(f"Unsupported mime type: {version.mime_type}")

            version.extracted_text = text
            version.processing_error = None
            document.status = DocumentStatus.completed
            db.commit()
            return {"status": "completed", "version_id": version_id}
        except Exception as e:
            version.processing_error = str(e)
            document.status = DocumentStatus.failed
            db.commit()
            return {"status": "failed", "version_id": version_id, "error": str(e)}
    finally:
        db.close()
